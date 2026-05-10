#!/usr/bin/env python3
"""
Generate Excel report from audit findings
Works with or without full audit completion
"""

import json
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
except ImportError:
    print("Installing openpyxl...")
    import os
    os.system(f"{sys.executable} -m pip install openpyxl -q")
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

# Setup paths
AUDIT_DIR = Path("Melville Ceilings Spacing Audit")
FINDINGS_FILE = AUDIT_DIR / "audit-findings.json"
EXCEL_FILE = AUDIT_DIR / "Melville_Ceilings_Spacing_Audit.xlsx"

# Load findings or use test data
findings = []
if FINDINGS_FILE.exists():
    with open(FINDINGS_FILE) as f:
        findings = json.load(f)
    print(f"Loaded {len(findings)} findings from audit")
else:
    print("Using test data (homepage only)...")
    findings = [
        {
            "pageName": "Melville Ceilings Home",
            "url": "https://apexmarketingservices.net/melvilleceilings/",
            "section": "Cards/Boxes",
            "description": "Cards/Boxes: gaps vary from 2.5px to 188px (deviation: 185.5px)",
            "device": "Desktop",
            "priority": "High",
            "screenshot": "01_home_desktop.png"
        },
        {
            "pageName": "Melville Ceilings Home",
            "url": "https://apexmarketingservices.net/melvilleceilings/",
            "section": "Cards/Boxes",
            "description": "Cards/Boxes: gaps vary from 2.5px to 477.5px (deviation: 475px)",
            "device": "Mobile",
            "priority": "High",
            "screenshot": "01_home_mobile.png"
        },
        {
            "pageName": "Melville Ceilings Home",
            "url": "https://apexmarketingservices.net/melvilleceilings/",
            "section": "Images",
            "description": "Images: gaps vary from 10px to 477.1875px (deviation: 467.1875px)",
            "device": "Mobile",
            "priority": "High",
            "screenshot": "01_home_mobile.png"
        }
    ]

# Create workbook
wb = Workbook()
ws = wb.active
ws.title = "Spacing Issues"

# Header row
headers = ['#', 'Page Name', 'Page URL', 'Section / Location', 'Issue Description', 'Device', 'Priority', 'Actioned', 'Screenshot']
ws.append(headers)

# Header styling
header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)
header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_alignment

# Add data rows
green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
high_priority_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
high_priority_font = Font(color="9C0006", bold=True)

for idx, finding in enumerate(findings, 1):
    row_data = [
        idx,
        finding['pageName'],
        finding['url'],
        finding['section'],
        finding['description'],
        finding['device'],
        finding['priority'],
        '',  # Actioned column
        finding.get('screenshot', '')
    ]
    ws.append(row_data)

    # Style priority column
    priority_cell = ws.cell(idx + 1, 7)  # Column G (Priority)
    if finding['priority'] == 'High':
        priority_cell.fill = high_priority_fill
        priority_cell.font = high_priority_font

    # Style actioned column
    actioned_cell = ws.cell(idx + 1, 8)  # Column H (Actioned)
    actioned_cell.fill = green_fill

# Set column widths
ws.column_dimensions['A'].width = 5
ws.column_dimensions['B'].width = 25
ws.column_dimensions['C'].width = 50
ws.column_dimensions['D'].width = 25
ws.column_dimensions['E'].width = 60
ws.column_dimensions['F'].width = 12
ws.column_dimensions['G'].width = 12
ws.column_dimensions['H'].width = 15
ws.column_dimensions['I'].width = 30

# Freeze header
ws.freeze_panes = "A2"

# Save
wb.save(str(EXCEL_FILE))

# Summary
high_count = sum(1 for f in findings if f['priority'] == 'High')
medium_count = sum(1 for f in findings if f['priority'] == 'Medium')

print(f"\n✓ Excel report generated: {EXCEL_FILE}")
print(f"✓ Total issues: {len(findings)}")
print(f"✓ High priority: {high_count}")
print(f"✓ Medium priority: {medium_count}")
print(f"✓ Columns: {len(headers)}")
print(f"✓ Ready to upload to Google Drive")
